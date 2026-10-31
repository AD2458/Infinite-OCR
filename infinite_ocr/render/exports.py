"""Export the IR to DOCX, XLSX, CSV, flattened Markdown and plain text.

All of these are pure post-processing over cached data, so every export costs
zero API credits no matter how many times it is regenerated.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from ..ir.models import Document
from ..script.detect import RTL_SCRIPTS
from ..stitch.levels import join_blocks
from .sanitize import strip_tags
from .tables import document_tables, grid_from_html, to_markdown_table
from .unwrap import unwrap_lines


_TABLE_RE = re.compile(r"<table\b.*?</table\s*>", re.IGNORECASE | re.DOTALL)


def _tables_to_markdown(text: str) -> str:
    """Replace every <table> in a fragment with a Markdown pipe table.

    Handles blocks that mix prose and table markup, and leaves any surrounding
    text intact rather than discarding it.
    """
    out: list[str] = []
    last = 0
    for m in _TABLE_RE.finditer(text):
        before = strip_tags(text[last : m.start()]).strip()
        if before:
            out.append(before)
        grid = grid_from_html(m.group(0))
        out.append(to_markdown_table(grid) if grid else strip_tags(m.group(0)))
        last = m.end()
    tail = strip_tags(text[last:]).strip()
    if tail:
        out.append(tail)
    return "\n\n".join(p for p in out if p)


def _merged_blocks(doc: Document):
    """Yield (block, text) with cross-page links already applied."""
    links = doc.accepted_links()
    consumed: set[str] = set()
    for block in doc.flow_blocks():
        if block.id in consumed:
            continue
        text = block.text
        cursor = block
        while (link := links.get(cursor.id)) is not None:
            nxt = doc.by_id(link.to_block)
            if nxt is None or nxt.id in consumed:
                break
            text = join_blocks(link, text, nxt.text)
            consumed.add(nxt.id)
            cursor = nxt
        yield block, text


# ----------------------------------------------------------------- markdown
def render_markdown_flat(doc: Document) -> str:
    """Markdown with HTML tables converted to pipe tables.

    Lossy: merged cells are flattened. That is exactly why `html` remains the
    default table format -- this export trades fidelity for readability.
    """
    parts: list[str] = []
    for block, text in _merged_blocks(doc):
        text = text.strip()
        if not text:
            continue
        # Key on the presence of a table, not the declared type: the model
        # returns table markup inside blocks typed Text.
        if "<table" in text.lower():
            parts.append(_tables_to_markdown(text))
        elif block.type == "Title":
            parts.append("# " + text.replace("\n", " "))
        elif block.type == "SectionHeader":
            parts.append("## " + text.replace("\n", " "))
        elif block.type == "Caption":
            parts.append("### " + text.replace("\n", " "))
        elif block.type == "Footnote":
            parts.append("> " + unwrap_lines(text).replace("\n", "\n> "))
        else:
            # Rejoin column-edge wraps so sentences aren't split mid-clause.
            parts.append(unwrap_lines(text, block.primary_script))
    return "\n\n".join(parts).strip() + "\n"


# --------------------------------------------------------------- plain text
def render_text(doc: Document) -> str:
    """All markup stripped. For search indexing and downstream NLP."""
    parts: list[str] = []
    for block, text in _merged_blocks(doc):
        clean = strip_tags(text).strip() if "<" in text else text.strip()
        clean = unwrap_lines(clean, block.primary_script)
        clean = re.sub(r"\n{3,}", "\n\n", clean)
        if clean:
            parts.append(clean)
    return "\n\n".join(parts).strip() + "\n"


def csv_bytes(grid: list[list[str]]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(grid)
    return buf.getvalue().encode("utf-8-sig")


def write_csvs(doc: Document, out_dir: str | Path, stem: str = "table") -> list[Path]:
    """Each table as its own UTF-8-BOM CSV (so Excel opens Indic text right)."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for i, (bid, grid) in enumerate(document_tables(doc), start=1):
        p = out_dir / f"{stem}_{i:02d}_{bid}.csv"
        with p.open("w", encoding="utf-8-sig", newline="") as fh:
            csv.writer(fh).writerows(grid)
        written.append(p)
    return written


# -------------------------------------------------------------------- docx
def write_docx(doc: Document, path: str | Path) -> Path:
    """Word document with real tables. Opens in Word/LibreOffice."""
    from docx import Document as Docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    out = Docx()
    for block, text in _merged_blocks(doc):
        text = text.strip()
        if not text:
            continue
        rtl = block.primary_script in RTL_SCRIPTS

        if block.type == "Table":
            grid = grid_from_html(text) if "<" in text else []
            if grid:
                width = max(len(r) for r in grid)
                t = out.add_table(rows=0, cols=width)
                t.style = "Table Grid"
                for r in grid:
                    cells = t.add_row().cells
                    for i, val in enumerate(r[:width]):
                        cells[i].text = val
                out.add_paragraph("")
                continue
            text = strip_tags(text)

        if block.type == "Title":
            p = out.add_heading(text.replace("\n", " "), level=1)
        elif block.type == "SectionHeader":
            p = out.add_heading(text.replace("\n", " "), level=2)
        elif block.type == "Caption":
            p = out.add_heading(text.replace("\n", " "), level=3)
        elif block.type == "List":
            for line in text.splitlines():
                if line.strip():
                    out.add_paragraph(line.strip(), style="List Bullet")
            continue
        else:
            p = out.add_paragraph(text)

        if rtl:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    out.save(str(path))
    return path


# -------------------------------------------------------------------- xlsx
def write_xlsx(doc: Document, path: str | Path) -> Path:
    """One sheet per table. Usually the most useful export for tabular PDFs."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    tables = document_tables(doc)

    if not tables:
        ws = wb.create_sheet("no_tables")
        ws["A1"] = "This document contained no tables."
    for i, (bid, grid) in enumerate(tables, start=1):
        ws = wb.create_sheet(f"{i}_{bid}"[:31])
        for r in grid:
            ws.append(r)
        if ws.max_row:
            head = PatternFill("solid", fgColor="EDEBFD")
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = head
            for col in ws.columns:
                longest = max((len(str(c.value or "")) for c in col), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(
                    60, max(12, longest + 2)
                )
            for row in ws.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path
