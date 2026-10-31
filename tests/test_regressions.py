"""Regression tests pinned to a REAL Bodhan API response.

Fixtures in tests/fixtures/rbi_page*.json are verbatim responses from a live
run on an RBI tender PDF. They encode behaviour the published docs did not
show -- `order` starting at 0, and types like `PageNumber`/`Folio` -- each of
which caused a real bug. Keeping the raw responses means these can never
silently regress.
"""

from __future__ import annotations

import collections
import json
from pathlib import Path

import pytest

from infinite_ocr.ir.furniture import mark_furniture
from infinite_ocr.ir.models import Document, Page
from infinite_ocr.ir.normalize import blocks_from_response
from infinite_ocr.render.markdown import render
from infinite_ocr.review.confidence import flag_blocks
from infinite_ocr.stitch.links import apply_links

FIX = Path(__file__).parent / "fixtures"
PAGE_H = 2992
PAGE_W = 2116


def _payloads():
    return [
        json.loads((FIX / f"rbi_page{i}.json").read_text(encoding="utf-8"))
        for i in range(3)
    ]


@pytest.fixture
def rbi_doc():
    doc = Document(doc_id="rbi", source_path="rbi.pdf")
    for i, payload in enumerate(_payloads()):
        doc.pages.append(
            Page(index=i, width=PAGE_W, height=PAGE_H, sha256=f"{i:064d}", dpi=300)
        )
        doc.blocks.extend(blocks_from_response(payload, doc_id="rbi", page=i))
    mark_furniture(doc.blocks, {i: PAGE_H for i in range(3)})
    apply_links(doc)
    flag_blocks(doc)
    return doc


class TestLiveApiShape:
    """The published example and the live API disagree; handle both."""

    def test_order_starts_at_zero_in_practice(self):
        orders = [b.get("order") for b in _payloads()[0]["blocks"]]
        assert min(orders) == 0, "fixture should capture 0-based ordering"

    def test_live_types_are_recognised(self):
        types = {
            b.type for b in blocks_from_response(_payloads()[0], doc_id="d", page=0)
        }
        # PageNumber/Folio were previously falling through to Unknown.
        assert "PageNumber" in types
        assert "Unknown" not in types


class TestBlockIdsAreUnique:
    """A collision is silently dangerous: a redo could rewrite the wrong text."""

    def test_no_duplicate_ids_across_document(self, rbi_doc):
        ids = [b.id for b in rbi_doc.blocks]
        dupes = {k: v for k, v in collections.Counter(ids).items() if v > 1}
        assert not dupes, f"duplicate block ids: {dupes}"

    def test_by_id_resolves_every_block(self, rbi_doc):
        for b in rbi_doc.blocks:
            assert rbi_doc.by_id(b.id) is b

    def test_zero_based_order_does_not_collide(self):
        payload = {
            "blocks": [
                {"order": 0, "type": "Text", "bbox_xyxy": [0, 0, 10, 10], "text": "a"},
                {"order": 1, "type": "Text", "bbox_xyxy": [0, 20, 10, 30], "text": "b"},
            ]
        }
        ids = [b.id for b in blocks_from_response(payload, doc_id="d", page=0)]
        assert len(set(ids)) == 2

    def test_one_based_order_also_works(self):
        payload = {
            "blocks": [
                {"order": 1, "type": "Text", "bbox_xyxy": [0, 0, 10, 10], "text": "a"},
                {"order": 2, "type": "Text", "bbox_xyxy": [0, 20, 10, 30], "text": "b"},
            ]
        }
        blocks = blocks_from_response(payload, doc_id="d", page=0)
        assert [b.text for b in blocks] == ["a", "b"]
        assert len({b.id for b in blocks}) == 2


class TestFurnitureOnRealData:
    def test_page_numbers_never_reach_the_document(self, rbi_doc):
        out = render(rbi_doc, with_anchors=False, bidi=False)
        for marker in ("7 | Page", "8 | Page", "9 | Page"):
            assert marker not in out

    def test_repeated_running_title_is_removed(self, rbi_doc):
        out = render(rbi_doc, with_anchors=False, bidi=False)
        assert "Comprehensive Annual Maintenance Service" not in out

    def test_ocr_noise_does_not_defeat_matching(self, rbi_doc):
        # One page produced "8 | Page\nha" -- exact matching missed it.
        noisy = [b for b in rbi_doc.blocks if "ha" in b.text and "Page" in b.text]
        assert noisy and all(b.is_furniture for b in noisy)

    def test_body_content_survives(self, rbi_doc):
        out = render(rbi_doc, with_anchors=False, bidi=False)
        assert "Schedule of Tender" in out
        assert "Important instructions" in out
        assert "<table" in out


class TestExportsOnRealData:
    def test_html_export_renders_tables(self, rbi_doc):
        from infinite_ocr.render.html import render_html

        html = render_html(rbi_doc, title="RBI")
        assert "<table" in html and "</html>" in html
        assert "7 | Page" not in html

    def test_flat_markdown_converts_tables(self, rbi_doc):
        from infinite_ocr.render.exports import render_markdown_flat

        md = render_markdown_flat(rbi_doc)
        assert "| ---" in md
        assert "<table" not in md

    def test_plain_text_has_no_markup(self, rbi_doc):
        from infinite_ocr.render.exports import render_text

        txt = render_text(rbi_doc)
        assert "<" not in txt
        assert "Schedule of Tender" in txt

    def test_docx_has_real_tables(self, rbi_doc, tmp_path):
        from docx import Document as Docx

        from infinite_ocr.render.exports import write_docx

        p = write_docx(rbi_doc, tmp_path / "o.docx")
        assert len(Docx(str(p)).tables) >= 1

    def test_xlsx_one_sheet_per_table(self, rbi_doc, tmp_path):
        from openpyxl import load_workbook

        from infinite_ocr.render.exports import write_xlsx
        from infinite_ocr.render.tables import document_tables

        p = write_xlsx(rbi_doc, tmp_path / "o.xlsx")
        assert len(load_workbook(p).sheetnames) == len(document_tables(rbi_doc))

    def test_csv_per_table(self, rbi_doc, tmp_path):
        from infinite_ocr.render.exports import write_csvs

        files = write_csvs(rbi_doc, tmp_path)
        assert files and all(f.stat().st_size > 0 for f in files)


class TestGridExpansion:
    def test_colspan_text_in_anchor_only(self):
        from infinite_ocr.render.tables import grid_from_html

        grid = grid_from_html(
            "<table><tr><td colspan='2'>wide</td></tr>"
            "<tr><td>a</td><td>b</td></tr></table>"
        )
        # Text lives in the anchor cell; spanned positions are empty.
        assert grid[0] == ["wide", ""]
        assert grid[1] == ["a", "b"]

    def test_rowspan_text_in_anchor_only(self):
        from infinite_ocr.render.tables import grid_from_html

        grid = grid_from_html(
            "<table><tr><td rowspan='2'>tall</td><td>x</td></tr>"
            "<tr><td>y</td></tr></table>"
        )
        assert grid[0][0] == "tall" and grid[1][0] == ""
        assert grid[0][1] == "x" and grid[1][1] == "y"

    def test_br_becomes_newline_not_stripped(self):
        from infinite_ocr.render.tables import grid_from_html

        grid = grid_from_html(
            "<table><tr><td>Sector<br/>துறை</td></tr></table>"
        )
        assert grid[0][0] == "Sector\nதுறை"

    def test_bilingual_cell_with_colspan(self):
        """The exact MSME regression: bilingual header with colspan=4."""
        from infinite_ocr.render.tables import grid_from_html

        grid = grid_from_html(
            "<table><tr><th colspan='4'>MSME Details<br/>MSME விவரங்கள்</th></tr>"
            "<tr><td>a</td><td>b</td><td>c</td><td>d</td></tr></table>"
        )
        assert grid[0][0] == "MSME Details\nMSME விவரங்கள்"
        assert grid[0][1:] == ["", "", ""]
        assert grid[1] == ["a", "b", "c", "d"]


class TestSanitizer:
    @pytest.mark.parametrize(
        "payload",
        [
            "<script>alert(1)</script>",
            "<img src=x onerror=alert(1)>",
            "<a href='javascript:alert(1)'>x</a>",
            "<div style='position:fixed'>x</div>",
            "<iframe src='http://evil'></iframe>",
        ],
    )
    def test_dangerous_markup_is_neutralised(self, payload):
        from infinite_ocr.render.sanitize import sanitize_html

        out = sanitize_html(payload)
        for bad in ("<script", "onerror", "javascript:", "<iframe", "style="):
            assert bad not in out

    def test_table_structure_is_preserved(self):
        from infinite_ocr.render.sanitize import sanitize_html

        out = sanitize_html("<table><tr><td colspan='3'>keep</td></tr></table>")
        assert "colspan" in out and "keep" in out

    def test_rtl_direction_attribute_survives(self):
        from infinite_ocr.render.sanitize import sanitize_html

        assert 'dir="rtl"' in sanitize_html('<p dir="rtl">اردو</p>')
